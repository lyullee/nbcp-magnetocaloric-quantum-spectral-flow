"""Build visually consistent Word submission files from the frozen Markdown.

The converter is intentionally local and deterministic.  It preserves the
manuscript wording, inserts all linked figures, turns mathematical spans into
Word OMML math objects, and formats Markdown tables for journal review.
"""

from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MANUSCRIPT = ROOT / "manuscript" / "Manuscript_Cryogenics_20261012.md"
SUPPLEMENT = ROOT / "manuscript" / "Supplementary_Information_Cryogenics_20261012.md"
MAIN_DOCX = ROOT / "manuscript" / "Manuscript_Cryogenics_20261012.docx"
SI_DOCX = ROOT / "manuscript" / "Supplementary_Information_Cryogenics_20261012.docx"

FONT = "Times New Roman"
MATH_FONT = "Cambria Math"
ACCENT = "17365D"
PALE = "EAF1F8"
BORDER = "D9D9D9"

SUPERSCRIPT = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")
SUBSCRIPT = str.maketrans("0123456789+-.()", "₀₁₂₃₄₅₆₇₈₉₊₋.₍₎")


def set_run_font(run, name=FONT, size=Pt(10.5), bold=None, italic=None):
    run.font.name = name
    run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def set_cell_margins(cell, top=85, start=90, bottom=85, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=Pt(9))


def normalize_plain_text(text):
    replacements = {
        "Na2BaCo(PO4)2": "Na₂BaCo(PO₄)₂",
        "Gd2B2MoO9": "Gd₂B₂MoO₉",
        "YbNi1.6Sn": "YbNi₁.₆Sn",
        "RuO2": "RuO₂",
        r"\times": "×",
        r"\mu": "μ",
        r"\Delta": "Δ",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"\(\^\{([0-9+-]+)\}\)", lambda m: m.group(1).translate(SUPERSCRIPT), text)
    text = re.sub(r"\^\{([0-9+-]+)\}", lambda m: m.group(1).translate(SUPERSCRIPT), text)
    text = re.sub(r"_\{([0-9+-.()]+)\}", lambda m: m.group(1).translate(SUBSCRIPT), text)
    text = re.sub(r"_([0-9])", lambda m: m.group(1).translate(SUBSCRIPT), text)
    return text


LATEX_MAP = {
    r"\Delta": "Δ", r"\delta": "δ", r"\theta": "θ", r"\lambda": "λ",
    r"\eta": "η", r"\rho": "ρ", r"\alpha": "α", r"\beta": "β",
    r"\mu": "μ", r"\sigma": "σ", r"\pi": "π", r"\epsilon": "ε",
    r"\Omega": "Ω", r"\sum": "Σ", r"\int": "∫", r"\infty": "∞",
    r"\geq": "≥", r"\ge": "≥", r"\leq": "≤", r"\le": "≤",
    r"\approx": "≈", r"\neq": "≠", r"\pm": "±", r"\mp": "∓",
    r"\times": "×", r"\cdot": "·", r"\rightarrow": "→", r"\to": "→",
    r"\parallel": "∥", r"\perp": "⊥", r"\in": "∈", r"\propto": "∝",
    r"\langle": "⟨", r"\rangle": "⟩", r"\left": "", r"\right": "",
    r"\qquad": "   ", r"\quad": "  ", r"\,": " ", r"\;": " ",
    r"\!": "", r"\%": "%", r"\bar": "",
}


def _replace_command_group(text, command):
    pattern = re.compile(r"\\" + command + r"\{([^{}]*)\}")
    previous = None
    while text != previous:
        previous = text
        text = pattern.sub(lambda m: m.group(1), text)
    return text


def clean_latex(value):
    """Convert the project's compact TeX vocabulary to readable linear math."""
    text = value.strip().replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\\begin\{(?:aligned|array|cases)\}", "", text)
    text = re.sub(r"\\end\{(?:aligned|array|cases)\}", "", text)
    text = text.replace(r"\\", " ; ")
    for cmd in ("mathrm", "textrm", "text", "rm", "operatorname"):
        text = _replace_command_group(text, cmd)
    # Flatten simple fractions and roots while retaining explicit grouping.
    frac = re.compile(r"\\frac\{([^{}]*)\}\{([^{}]*)\}")
    previous = None
    while text != previous:
        previous = text
        text = frac.sub(lambda m: f"({m.group(1)})/({m.group(2)})", text)
    text = re.sub(r"\\sqrt\{([^{}]*)\}", lambda m: f"√({m.group(1)})", text)
    for src, dst in LATEX_MAP.items():
        text = text.replace(src, dst)
    text = re.sub(r"\\(arg|max|min|median|exp|sin|cos|log|ln)\b", lambda m: m.group(1), text)
    text = re.sub(r"_\{([^{}]+)\}", lambda m: "_(" + m.group(1) + ")", text)
    text = re.sub(r"\^\{([^{}]+)\}", lambda m: "^(" + m.group(1) + ")", text)
    text = re.sub(r"_([A-Za-z0-9])", r"_(\1)", text)
    text = re.sub(r"\^([A-Za-z0-9*+-])", r"^(\1)", text)
    text = text.replace("{", "(").replace("}", ")")
    text = re.sub(r"\\([A-Za-z]+)", lambda m: m.group(1), text)
    text = text.replace("&", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def append_math(paragraph, latex, display=False):
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    rpr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    rpr.append(sty)
    run.append(rpr)
    text = OxmlElement("m:t")
    text.text = clean_latex(latex)
    run.append(text)
    math.append(run)
    if display:
        math_para = OxmlElement("m:oMathPara")
        math_para.append(math)
        paragraph._p.append(math_para)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph._p.append(math)


TOKEN_RE = re.compile(
    r"(\\\(.*?\\\)|\$[^$\n]+\$|\*\*.*?\*\*|`[^`]+`|\*[^*\n]+\*|<sup>.*?</sup>|<sub>.*?</sub>)"
)


def add_inline(paragraph, text, base_size=Pt(10.5), force_bold=False, force_italic=False):
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(normalize_plain_text(text[pos:match.start()]))
            set_run_font(run, size=base_size, bold=force_bold, italic=force_italic)
        token = match.group(0)
        if token.startswith(r"\("):
            append_math(paragraph, token[2:-2], display=False)
        elif token.startswith("$"):
            append_math(paragraph, token[1:-1], display=False)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, italic=force_italic)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=Pt(max(8.5, base_size.pt - 1)))
        elif token.startswith("<sup>"):
            run = paragraph.add_run(token[5:-6])
            set_run_font(run, size=Pt(base_size.pt - 1))
            run.font.superscript = True
        elif token.startswith("<sub>"):
            run = paragraph.add_run(token[5:-6])
            set_run_font(run, size=Pt(base_size.pt - 1))
            run.font.subscript = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, bold=force_bold, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(normalize_plain_text(text[pos:]))
        set_run_font(run, size=base_size, bold=force_bold, italic=force_italic)


def strip_markup(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = text.replace("`", "")
    text = re.sub(r"</?(?:sup|sub)>", "", text)
    return text


def configure_document(doc, short_title, supplement=False):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.70)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.28)
    sec.footer_distance = Inches(0.30)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(17)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    heading_specs = {
        "Heading 1": (14, 12, 6),
        "Heading 2": (12, 9, 4),
        "Heading 3": (11, 7, 3),
    }
    for name, (size, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption" in styles:
        caption = styles["Caption"]
        caption.font.name = FONT
        caption.font.size = Pt(9)
        caption.font.italic = False
        caption.font.color.rgb = RGBColor(0, 0, 0)
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(7)
        caption.paragraph_format.keep_together = True

    header = sec.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(("Supplementary Information  |  " if supplement else "Research article  |  ") + short_title)
    set_run_font(run, size=Pt(8.5))
    run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_field(sec.footer.paragraphs[0])

    doc.core_properties.title = short_title
    doc.core_properties.subject = "Cryogenics submission manuscript" if not supplement else "Cryogenics supplementary information"
    doc.core_properties.author = "Ugwiyeon Lee; Seong-Yun Bae; Hakrae Kim; Jungsu Lee"
    doc.core_properties.keywords = "magnetic refrigeration, quantum hardware, Na2BaCo(PO4)2, cryogenics"


def add_table(doc, rows):
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            value = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, strip_markup(value), base_size=Pt(8.1), force_bold=(r_idx == 0))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, ACCENT)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, source_path, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    with Image.open(source_path) as img:
        width_px, height_px = img.size
    max_w, max_h = 6.75, 7.25
    aspect = width_px / max(height_px, 1)
    width = min(max_w, max_h * aspect)
    height = width / aspect
    if height > max_h:
        height = max_h
        width = height * aspect
    p.add_run().add_picture(str(source_path), width=Inches(width), height=Inches(height))
    # Alternative text in docPr for accessibility.
    drawing = p._p.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", alt)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell or "---") for cell in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build(markdown_path, output_path, supplement=False):
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    short_title = "NBCP magnetic refrigerator stage"
    configure_document(doc, short_title, supplement=supplement)

    i = 0
    in_code = False
    code_lines = []
    first_title = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(6)
                for n, code in enumerate(code_lines):
                    if n:
                        p.add_run().add_break()
                    run = p.add_run(code)
                    set_run_font(run, name="Consolas", size=Pt(8.5))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line == r"\[":
            equation = []
            i += 1
            while i < len(lines) and lines[i].strip() != r"\]":
                equation.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_together = True
            append_math(p, " ".join(equation), display=True)
            i += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            image_path = (markdown_path.parent / image_match.group(2)).resolve()
            add_figure(doc, image_path, image_match.group(1))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[i + 1].strip()):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline(p, line[2:], base_size=Pt(17), force_bold=True)
            first_title = False
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:], base_size=Pt(14), force_bold=True)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], base_size=Pt(12), force_bold=True)
        elif line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[5:], base_size=Pt(11), force_bold=True)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2.5)
            add_inline(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            # Preserve manuscript reference numbers rather than Word auto-numbering.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line, base_size=Pt(9.3))
        elif line.startswith("**Figure") or line.startswith("**Table"):
            p = doc.add_paragraph(style="Caption")
            add_inline(p, line, base_size=Pt(9))
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            add_inline(p, line[2:-2], force_bold=True)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if markdown_path == MANUSCRIPT and i in (2, 4, 6):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if "Cryogenics-targeted working draft" in line:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
            add_inline(p, line)
        i += 1

    doc.save(output_path)
    return output_path


def main():
    outputs = [
        build(MANUSCRIPT, MAIN_DOCX, supplement=False),
        build(SUPPLEMENT, SI_DOCX, supplement=True),
    ]
    if sys.platform == "win32":
        formatter = ROOT / "professionalize_docx_math.ps1"
        for path in outputs:
            subprocess.run(
                [
                    "powershell",
                    "-NoProfile",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    str(formatter),
                    "-InputDocx",
                    str(path),
                ],
                check=True,
            )
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
