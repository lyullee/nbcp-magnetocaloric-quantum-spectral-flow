"""Audit the published refrigeration-comparison source workbook.

The source workbook is the public Source Data for Extended Data Fig. 6 of
Xiang et al., Nature 625, 270-275 (2024). Raw workbook values remain
unchanged. The workbook labels in sheets a/b do not reproduce the article's
stated 94 mK NBCP point at about 1.5 T: that point occurs on the curve labelled
GGG-2-4. We therefore expose the mismatch and use only the independently
labelled sheet-d time traces for a quantitative cross-material summary.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "sources" / "extended_figure_6.xlsx"
RESULT_CSV = ROOT / "results" / "refrigerant_experimental_benchmark.csv"
RESULT_JSON = ROOT / "results" / "refrigerant_experimental_benchmark.json"

SAMPLE_META = {
    "NBCP": {"mass_g": 1.65, "volume_cm3": 0.395},
    "CPA": {"mass_g": 0.76, "volume_cm3": 0.421},
    "FAA": {"mass_g": 0.70, "volume_cm3": 0.409},
}


def series(sheet, x_col: int, y_col: int) -> list[tuple[float, float]]:
    points: list[tuple[float, float]] = []
    for row in range(6, sheet.max_row + 1):
        x = sheet.cell(row, x_col).value
        y = sheet.cell(row, y_col).value
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            points.append((float(x), float(y)))
    return points


def label_material(label: str) -> str:
    for material in SAMPLE_META:
        if material in label:
            return material
    raise ValueError(f"Unrecognized material label: {label!r}")


def threshold_span(points: list[tuple[float, float]], threshold: float) -> float | None:
    below = [x for x, temperature in points if temperature < threshold]
    return max(below) - min(below) if below else None


def nearest_temperature(points: list[tuple[float, float]], x_target: float) -> tuple[float, float]:
    return min(points, key=lambda point: abs(point[0] - x_target))


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)

    # Preserve sheet-b extrema exactly as labelled, but do not use them for a
    # material ranking because the sheet-a/b labels conflict with the article
    # caption at the identifiable 94 mK, B about 1.5 T point.
    demag = workbook["b"]
    raw_extrema: list[dict[str, float | str]] = []
    for x_col in (1, 3, 5):
        label = str(demag.cell(5, x_col + 1).value)
        material = label_material(label)
        points = series(demag, x_col, x_col + 1)
        field_drop_at_min, minimum_temperature = min(points, key=lambda point: point[1])
        raw_extrema.append(
            {
                "material": material,
                "source_curve": label,
                "mass_g": SAMPLE_META[material]["mass_g"],
                "volume_cm3": SAMPLE_META[material]["volume_cm3"],
                "minimum_temperature_K": minimum_temperature,
                "field_drop_at_minimum_T": field_drop_at_min,
                "field_at_minimum_T": 4.0 - field_drop_at_min,
                "curve_points": len(points),
            }
        )

    # Locate the article's stated NBCP benchmark in sheet a. For B0=4 T and a
    # remaining field of 1.5 T, the workbook abscissa B0-B is 2.5 T.
    sheet_a = workbook["a"]
    caption_check = []
    for x_col in (1, 3, 5, 8, 10, 12):
        label = str(sheet_a.cell(5, x_col + 1).value)
        x_value, temperature = nearest_temperature(series(sheet_a, x_col, x_col + 1), 2.5)
        caption_check.append({
            "workbook_label": label,
            "sampled_field_drop_T": x_value,
            "remaining_field_T_assuming_B0_4T": 4.0 - x_value,
            "temperature_K": temperature,
        })
    closest_to_reported_94mK = min(
        caption_check, key=lambda row: abs(row["temperature_K"] - 0.094)
    )

    # Independent time traces (Extended Data Fig. 6d). This span is the
    # observed interval between first and last samples below 1 K; it is not a
    # fitted thermal time constant.
    time_sheet = workbook["d"]
    rows: list[dict[str, float | str | None]] = []
    for x_col in (1, 3, 5):
        label = str(time_sheet.cell(5, x_col + 1).value)
        material = label_material(label)
        points = series(time_sheet, x_col, x_col + 1)
        rows.append({
            "material": material,
            "source_curve": label,
            "mass_g": SAMPLE_META[material]["mass_g"],
            "volume_cm3": SAMPLE_META[material]["volume_cm3"],
            "below_1K_observed_span_min": threshold_span(points, 1.0),
            "minimum_temperature_K": min(temperature for _, temperature in points),
            "time_trace_points": len(points),
        })

    RESULT_CSV.parent.mkdir(exist_ok=True)
    with RESULT_CSV.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    summary = {
        "status": "published refrigeration source-data audit",
        "source": {
            "doi": "10.1038/s41586-023-06885-w",
            "workbook": "sources/extended_figure_6.xlsx",
            "sheet_a_b": "demagnetization curves with a label/caption consistency issue",
            "sheet_d": "temperature-time traces under the source experiment conditions",
        },
        "time_trace_rows_safe_for_comparison": rows,
        "sheet_b_raw_extrema_by_workbook_label_not_used_for_ranking": raw_extrema,
        "caption_consistency_check": {
            "article_statement": "NBCP reaches approximately 94 mK at about 1.5 T from B0=4 T, T0=2 K",
            "curve_closest_to_94mK_at_field_drop_about_2p5T": closest_to_reported_94mK,
            "all_sheet_a_curves_at_field_drop_about_2p5T": caption_check,
            "labels_consistent_with_article_statement": closest_to_reported_94mK["workbook_label"].startswith("NBCP"),
        },
        "interpretation": [
            "Sheet-a/b workbook labels do not reproduce the identifiable article benchmark and are not used for a material ranking.",
            "The below-1 K span is the time between the first and last sampled values below 1 K.",
            "The comparison establishes material promise but does not determine a robust terminal field, field tolerance, or full-device inventory.",
        ],
    }
    RESULT_JSON.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
